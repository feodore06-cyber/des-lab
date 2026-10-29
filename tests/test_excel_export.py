# -*- coding: utf-8 -*-
"""The clean Excel export: sheets, formatting and a lossless round trip."""
import io
import os

os.environ["DESLAB_STORAGE"] = "file"  # tests must never touch a real database
os.environ["DESLAB_MAIL_DISABLE"] = "1"  # tests must never send mail

import sys
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import excel_export  # noqa: E402
import excel_io  # noqa: E402

FRAME = pd.DataFrame([
    {"sample_id": "EXP-01", "feedstock": "wood flour", "des_hba": "ChCl",
     "des_hbd": "Glyoxylic acid", "molar_ratio": 2.0, "purity_wt_pct": 98.0,
     "water_wt_pct": 0.0, "biomass_g": 2.8, "temperature_C": 80.0, "time_min": 60.0,
     "lignin_mg": 56.0, "lignin_yield_pct": 2.0, "replicate": 1,
     "recorded_at": "2026-10-29 10:00", "_owner": "lab@example.com"},
])


class CleanExcelTests(unittest.TestCase):
    def setUp(self):
        self.payload = excel_export.records_workbook(FRAME, meta={"user": "lab@example.com"})
        self.book = load_workbook(io.BytesIO(self.payload))

    def test_the_workbook_has_a_records_sheet_and_a_guide(self):
        self.assertEqual(self.book.sheetnames[:2], ["Records", "Guide"])

    def test_the_header_carries_units_and_is_frozen(self):
        sheet = self.book["Records"]
        headers = [cell.value for cell in sheet[3]]
        self.assertIn("Biomass (g)", headers)
        self.assertIn("Molar ratio HBD:HBA", headers)
        self.assertEqual(sheet.freeze_panes, "A4", "the header row stays visible while scrolling")

    def test_numbers_stay_numbers(self):
        sheet = self.book["Records"]
        row = [cell.value for cell in sheet[4]]
        numbers = [value for value in row if isinstance(value, (int, float))]
        self.assertGreaterEqual(len(numbers), 5, "masses and ratios must not become text")

    def test_the_owner_column_is_renamed_for_humans(self):
        sheet = self.book["Records"]
        headers = [cell.value for cell in sheet[3]]
        self.assertIn("Recorded by", headers)

    def test_the_guide_explains_every_column(self):
        guide = self.book["Guide"]
        text = " ".join(str(cell.value) for row in guide.iter_rows() for cell in row if cell.value)
        self.assertIn("lignin_mg", text.replace(" ", "").lower().replace("lignin(mg)", "lignin_mg")
                      if False else text)
        self.assertIn("hydrogen-bond acceptor", text)
        self.assertIn("Exported", text)

    def test_the_exported_file_can_be_imported_back_unchanged(self):
        result = excel_io.read_table(self.payload)
        back = result["frame"]
        self.assertEqual(len(back), 1)
        self.assertEqual(back.loc[0, "sample_id"], "EXP-01")
        self.assertEqual(back.loc[0, "des_hba"], "ChCl")
        self.assertAlmostEqual(float(back.loc[0, "molar_ratio"]), 2.0)
        self.assertAlmostEqual(float(back.loc[0, "lignin_yield_pct"]), 2.0)

    def test_an_empty_table_still_produces_a_valid_file(self):
        empty = pd.DataFrame(columns=["sample_id", "biomass_g", "lignin_mg"])
        payload = excel_export.records_workbook(empty)
        book = load_workbook(io.BytesIO(payload))
        self.assertIn("Records", book.sheetnames)


if __name__ == "__main__":
    unittest.main(verbosity=2)
