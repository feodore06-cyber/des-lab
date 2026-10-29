# -*- coding: utf-8 -*-
"""Write the records as a clean, re-importable Excel workbook.

Three sheets, so somebody who has never seen the project can read it:

* ``Records``  - one row per experiment, canonical columns, units in the header, frozen header row
* ``Summary``  - the grouped statistics, when there are any
* ``Guide``    - what every column means, who exported it and when

The output is deliberately round-trip safe: ``excel_io.read_table`` reads it back unchanged.
"""
from __future__ import annotations

import io
import time

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="002F86")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(color="0B1F4B", bold=True, size=14)
NOTE_FONT = Font(color="55637D", size=9, italic=True)
THIN = Side(style="thin", color="DCE4F2")
BORDER = Border(bottom=THIN)

# column -> (header text, number format)
COLUMN_STYLE = {
    "sample_id": ("Sample ID", None),
    "formulation": ("Formulation (label)", None),
    "feedstock": ("Feedstock", None),
    "des_hba": ("H-bond acceptor", None),
    "des_hbd": ("H-bond donor", None),
    "molar_ratio": ("Molar ratio HBD:HBA", "0.###"),
    "purity_wt_pct": ("Purity (wt%)", "0.0"),
    "water_wt_pct": ("Water (wt%)", "0.0"),
    "biomass_g": ("Biomass (g)", "0.0000"),
    "temperature_C": ("Temperature (°C)", "0.0"),
    "time_min": ("Time (min)", "0"),
    "lignin_mg": ("Lignin (mg)", "0.000"),
    "lignin_yield_pct": ("Lignin yield (%)", "0.0000"),
    "replicate": ("Replicate", "0"),
    "recorded_at": ("Recorded at", None),
    "recorded_by": ("Recorded by", None),
    "notes": ("Notes", None),
}

COLUMN_HELP = {
    "sample_id": ("text", "your own identifier for the run, e.g. EXP-01"),
    "formulation": ("text", "human label built from the fields below, e.g. ChCl:Glyoxylic acid 1:2"),
    "feedstock": ("text", "what was pretreated, e.g. wood flour, rice straw"),
    "des_hba": ("text", "hydrogen-bond acceptor, canonical name (ChCl)"),
    "des_hbd": ("text", "hydrogen-bond donor, canonical name (Glyoxylic acid)"),
    "molar_ratio": ("number", "moles of donor per mole of acceptor: 1:2 is 2"),
    "purity_wt_pct": ("number", "reagent purity from the bottle label, percent"),
    "water_wt_pct": ("number", "water added on purpose, percent of the reagent mass"),
    "biomass_g": ("number", "dry biomass charged, grams"),
    "temperature_C": ("number", "pretreatment temperature, °C"),
    "time_min": ("number", "pretreatment time, minutes"),
    "lignin_mg": ("number", "lignin recovered, milligrams"),
    "lignin_yield_pct": ("number", "lignin_mg / (biomass_g × 1000) × 100"),
    "replicate": ("number", "which repeat of the same condition (1, 2, 3 …)"),
    "recorded_at": ("text", "when the row was written, local time"),
    "recorded_by": ("text", "the account that recorded the row"),
    "notes": ("text", "free text"),
}


def _write_table(sheet, frame: pd.DataFrame, title: str) -> None:
    sheet["A1"] = title
    sheet["A1"].font = TITLE_FONT
    sheet.row_dimensions[1].height = 22

    header_row = 3
    for index, column in enumerate(frame.columns, start=1):
        label = COLUMN_STYLE.get(column, (column, None))[0]
        cell = sheet.cell(row=header_row, column=index, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = BORDER
    sheet.row_dimensions[header_row].height = 28

    for row_index, (_, row) in enumerate(frame.iterrows(), start=header_row + 1):
        for column_index, column in enumerate(frame.columns, start=1):
            value = row[column]
            if isinstance(value, float) and pd.isna(value):
                value = None
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            number_format = COLUMN_STYLE.get(column, (None, None))[1]
            if number_format and value is not None:
                cell.number_format = number_format
            cell.alignment = Alignment(horizontal="left", vertical="top")

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    if len(frame):
        sheet.auto_filter.ref = (f"A{header_row}:"
                                 f"{get_column_letter(len(frame.columns))}{header_row + len(frame)}")
    for index, column in enumerate(frame.columns, start=1):
        widest = max([len(str(COLUMN_STYLE.get(column, (column, None))[0]))] +
                     [len(str(value)) for value in frame[column].head(200)])
        sheet.column_dimensions[get_column_letter(index)].width = min(34, max(11, widest + 2))


def records_workbook(frame: pd.DataFrame, summary: pd.DataFrame | None = None,
                     meta: dict | None = None) -> bytes:
    """Build the workbook and return its bytes."""
    workbook = Workbook()
    records = workbook.active
    records.title = "Records"
    clean = frame.copy()
    if "_owner" in clean.columns:
        clean = clean.rename(columns={"_owner": "recorded_by"})
    ordered = [column for column in COLUMN_STYLE if column in clean.columns]
    ordered += [column for column in clean.columns if column not in ordered]
    _write_table(records, clean[ordered], "Experiment records")

    if summary is not None and not summary.empty:
        sheet = workbook.create_sheet("Summary")
        _write_table(sheet, summary, "Grouped summary (n, mean, sample SD)")

    guide = workbook.create_sheet("Guide")
    guide["A1"] = "How to read this file"
    guide["A1"].font = TITLE_FONT
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 12
    guide.column_dimensions["C"].width = 92
    guide["A3"], guide["B3"], guide["C3"] = "Column", "Type", "Meaning"
    for cell in (guide["A3"], guide["B3"], guide["C3"]):
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    row_index = 4
    for column in ordered:
        kind, meaning = COLUMN_HELP.get(column, ("text", ""))
        guide.cell(row=row_index, column=1, value=COLUMN_STYLE.get(column, (column, None))[0])
        guide.cell(row=row_index, column=2, value=kind)
        guide.cell(row=row_index, column=3, value=meaning)
        row_index += 1

    row_index += 1
    guide.cell(row=row_index, column=1, value="Provenance").font = Font(bold=True)
    row_index += 1
    info = {
        "Exported": time.strftime("%Y-%m-%d %H:%M"),
        "Rows": len(clean),
        "Exported by": (meta or {}).get("user", ""),
        "Scope": (meta or {}).get("scope", ""),
        "Upload this file back": "Results -> My records -> Import from Excel",
    }
    for label, value in info.items():
        guide.cell(row=row_index, column=1, value=label)
        guide.cell(row=row_index, column=3, value=value)
        row_index += 1

    guide.cell(row=row_index + 1, column=1,
               value="The Records sheet can be uploaded back into BioMat Lab unchanged.").font = NOTE_FONT

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
